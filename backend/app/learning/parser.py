"""多格式文档解析：根据扩展名自动选择解析方式，统一抽取为纯文本。

支持：pdf / docx / xlsx / csv / html / json / md / txt 等；
未知扩展名按纯文本尝试（utf-8 -> gb18030 -> latin-1）。
"""

import csv
import io
import json

from bs4 import BeautifulSoup


def decode_bytes(data: bytes) -> str:
    for enc in ("utf-8", "gb18030", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text.strip())
    return "\n\n".join(pages)


def _parse_docx(data: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"# 工作表: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                lines.append(" | ".join(cells))
    wb.close()
    return "\n".join(lines)


def _parse_csv(data: bytes) -> str:
    text = decode_bytes(data)
    reader = csv.reader(io.StringIO(text))
    return "\n".join(" | ".join(row) for row in reader if row)


def _parse_json(data: bytes) -> str:
    obj = json.loads(decode_bytes(data))
    if isinstance(obj, list) and obj and all(isinstance(x, dict) for x in obj):
        # 记录列表：每条记录展开为 "字段: 值" 文本块，便于检索
        blocks = []
        for i, item in enumerate(obj, 1):
            lines = [f"[记录 {i}]"]
            lines += [f"{k}: {v}" for k, v in item.items()]
            blocks.append("\n".join(lines))
        return "\n\n".join(blocks)
    return json.dumps(obj, ensure_ascii=False, indent=2)


def parse_document(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return _parse_pdf(data)
    if name.endswith(".docx"):
        return _parse_docx(data)
    if name.endswith((".xlsx", ".xlsm")):
        return _parse_xlsx(data)
    if name.endswith(".csv"):
        return _parse_csv(data)
    if name.endswith((".html", ".htm")):
        soup = BeautifulSoup(decode_bytes(data), "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        return soup.get_text("\n")
    if name.endswith(".json"):
        return _parse_json(data)
    return decode_bytes(data)
